"""Layanan Pembangkitan Berkas Excel untuk Analisis Peringatan Dini & Rekapitulasi KPS."""

from __future__ import annotations

from datetime import datetime
import io
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from app.services.early_warning_service import EarlyWarningService


def export_early_warning_excel(
    service: EarlyWarningService,
    category: str = "all",
    wilker_bps: str | None = None,
    items: list[dict[str, Any]] | None = None,
    custom_title: str | None = None,
    custom_subtitle: str | None = None,
    province: str | None = None,
    skema: str | None = None,
    search: str | None = None,
    zone: str | None = None,
) -> bytes:
    """Bangun file Excel ekspor terstruktur langsung dari memori.

    Mendukung items yang sudah difilter/diurutkan di klien, atau query langsung
    ke database bila items tidak disediakan.
    """
    wb = openpyxl.Workbook()
    font_title = Font(name="Calibri", size=15, bold=True, color="1B365D")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=10)

    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_red_light = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    fill_orange_light = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    fill_yellow_light = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="D3D3D3"), right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"), bottom=Side(style="thin", color="D3D3D3")
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    ws = wb.active
    ws.title = "Rekap Analisis Kebakaran"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = custom_title or "REKAPITULASI ANALISIS KEBAKARAN & PERINGATAN DINI KPS"
    ws["A1"].font = font_title

    if items is None:
        items = service.get_kps_analysis_list(
            category=category,
            wilker_bps=wilker_bps,
            province=province,
            skema=skema,
            search=search,
            limit=2500,
        )
        if zone and zone != "all":
            items = [r for r in items if r.get("zone_code") == zone]

    if custom_subtitle:
        sub_text = f"Dihasilkan: {datetime.now().strftime('%d %B %Y %H:%M WIB')} | {custom_subtitle} | Total: {len(items)} KPS"
    else:
        wilker_str = f" | Balai PS: {wilker_bps}" if wilker_bps else ""
        sub_text = f"Dihasilkan pada: {datetime.now().strftime('%d %B %Y %H:%M WIB')} | Kategori: {category.upper()}{wilker_str} | Total: {len(items)} KPS"
    ws["A2"] = sub_text
    ws["A2"].font = font_subtitle

    headers = [
        "No", "ID", "Nama KPS / Lembaga", "Balai PS", "Skema", "Desa", "Kecamatan", "Kabupaten", "Provinsi",
        "Luas SK (ha)", "Luas Terbakar Kemenhut (ha)", "Frekuensi Terbakar",
        "Hotspot Hari Ini (Total)", "Tepat di Bekas Terbakar (Strict Re-burn)", "Perembetan Blok Baru",
        "Jarak Perambatan Min (KM)", "Jarak Perambatan Max (KM)", "Jarak Perambatan Rerata (KM)",
        "Arah Perambatan Kompas", "Sudut Azimuth (°)",
        "Klasifikasi Zona Ilmiah",
        "Hotspot Kemarin", "Hotspot 7 Hari", "Hotspot Bulan Ini", "Total Hotspot 2026",
        "Deteksi Terakhir", "Skor FTRI", "Status & Rincian Mekanisme"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    for row_idx, r in enumerate(items, 5):
        raw_dt = r.get("latest_hotspot_at")
        dt_str = str(raw_dt)[:16].replace("T", " ") if raw_dt else "-"

        luas_sk = r.get("luas_sk")
        luas_sk_val = float(luas_sk) if luas_sk is not None else None

        total_burned = r.get("total_burned_ha")
        total_burned_val = float(total_burned) if total_burned is not None else 0.0

        burn_freq = r.get("burn_frequency")
        burn_freq_val = int(burn_freq) if burn_freq is not None else 0

        h_today = int(r.get("hotspots_today") or 0)
        h_strict = int(r.get("hotspots_today_strict_reburn") or 0)
        h_expanding = int(r.get("hotspots_today_expanding") or 0)

        min_d = r.get("min_distance_km")
        max_d = r.get("max_distance_km")
        avg_d = r.get("avg_distance_km")

        h_yest = int(r.get("hotspots_yesterday") or 0)
        h_7d = int(r.get("hotspots_7d") or 0)
        h_month = int(r.get("hotspots_month") or 0)
        h_year = int(r.get("hotspots_year") or 0)

        ftri = float(r.get("ftri_score") or 0)
        zone_code = r.get("zone_code")

        row_values = [
            row_idx - 4,
            r.get("id"),
            r.get("lembaga"),
            r.get("wilker_bps") or "-",
            r.get("skema") or "-",
            r.get("nama_desa") or "-",
            r.get("nama_kec") or "-",
            r.get("nama_kab") or "-",
            r.get("nama_prov") or "-",
            luas_sk_val,
            total_burned_val,
            burn_freq_val,
            h_today,
            h_strict,
            h_expanding,
            min_d,
            max_d,
            avg_d,
            r.get("fire_direction") or "-",
            r.get("fire_azimuth_deg") or "-",
            r.get("propagation_zone") or "-",
            h_yest,
            h_7d,
            h_month,
            h_year,
            dt_str,
            ftri,
            r.get("status_label") or "-",
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = border_thin
            if h_today > 0:
                if col_idx in [1, 13, 14, 15, 21, 27, 28]:
                    if zone_code in ["zone1", "strict", "combo"]:
                        cell.fill = fill_red_light
                    elif zone_code == "zone2":
                        cell.fill = fill_orange_light
                    else:
                        cell.fill = fill_yellow_light
            elif h_yest > 0:
                if col_idx in [1, 22, 27, 28]:
                    cell.fill = fill_orange_light

            if col_idx in [1, 2, 4, 5, 12, 19, 21, 26, 28]:
                cell.alignment = align_center
            elif col_idx in [10, 11, 13, 14, 15, 16, 17, 18, 20, 22, 23, 24, 25, 27]:
                cell.alignment = align_right
                if col_idx in [10, 11, 16, 17, 18, 20, 27]:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.alignment = align_left

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row > 2:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

"""Service Analisis Siaga Rambatan Api (Deteksi Ancaman Hotspot Luar Poligon KPS).

Mendeteksi hotspot di luar batas kawasan KPS (radius buffer 1 km, 3 km, 5 km)
yang berpotensi merambat masuk ke dalam areal KPS sebelum kebakaran meluas.
"""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import get_settings
from app.services.cache_service import CacheService
from app.services.postgres_store import PostgresStore

logger = logging.getLogger("fire_spread.service")

# 1 derajat di ekuator ~ 111.320 m -> 1 meter ~ 0.0000089831 derajat
DEGREE_PER_METER = 1.0 / 111320.0


def degrees_to_compass(deg: float | None) -> str:
    """Mengubah derajat bearing (0 - 360) menjadi arah mata angin Indonesia."""
    if deg is None:
        return "Tidak diketahui"
    deg = (deg % 360 + 360) % 360
    directions = [
        ("Utara (N)", 0, 22.5),
        ("Timur Laut (NE)", 22.5, 67.5),
        ("Timur (E)", 67.5, 112.5),
        ("Tenggara (SE)", 112.5, 157.5),
        ("Selatan (S)", 157.5, 202.5),
        ("Barat Daya (SW)", 202.5, 247.5),
        ("Barat (W)", 247.5, 292.5),
        ("Barat Laut (NW)", 292.5, 337.5),
        ("Utara (N)", 337.5, 360),
    ]
    for label, lower, upper in directions:
        if lower <= deg <= upper:
            return label
    return "Utara (N)"


def distance_to_level(dist_m: int) -> tuple[str, str]:
    if dist_m < 1000:
        return "bahaya", "Bahaya Kritis (< 1 km)"
    if dist_m < 3000:
        return "waspada", "Waspada (1–3 km)"
    return "pantau", "Pantau (3–5 km)"


class FireSpreadService:
    def __init__(
        self,
        postgres_store: PostgresStore | None = None,
        cache_service: CacheService | None = None,
    ) -> None:
        settings = get_settings()
        self.postgres_store = postgres_store or PostgresStore(settings.database_url)
        self.cache_service = cache_service or CacheService(
            Path(settings.cache_dir),
            settings.cache_ttl_hours,
            settings.database_url,
        )

    def get_summary(
        self,
        time_window_hours: int = 48,
        max_distance_km: float = 5.0,
        province: str | None = None,
        regency: str | None = None,
        wilker: str | None = None,
    ) -> dict[str, Any]:
        """Ringkasan statistik KPS yang terancam api di perimeter luar."""
        cache_key = f"fire_spread_summary_{time_window_hours}_{max_distance_km:.2f}_{province or 'all'}_{regency or 'all'}_{wilker or 'all'}"
        cached = self.cache_service.read(cache_key)
        if cached is not None and isinstance(cached, dict):
            return cached

        if not self.postgres_store.enabled:
            return {
                "total_kps_threatened": 0,
                "bahaya_count": 0,
                "waspada_count": 0,
                "pantau_count": 0,
                "total_external_hotspots": 0,
                "time_window_hours": time_window_hours,
                "max_distance_km": max_distance_km,
            }

        max_deg = max_distance_km * 1000.0 * DEGREE_PER_METER
        where_clauses = [
            "poly.is_active = true",
            f"obs.detected_at >= NOW() - INTERVAL '{int(time_window_hours)} hours'",
            f"ST_DWithin(poly.geometry, obs.geom, {max_deg:.6f})",
            "NOT ST_Intersects(poly.geometry, obs.geom)",
        ]
        params: list[Any] = []
        if province:
            where_clauses.append("poly.nama_prov = %s")
            params.append(province)
        if regency:
            where_clauses.append("poly.nama_kab = %s")
            params.append(regency)
        if wilker:
            where_clauses.append("poly.wilker_bps = %s")
            params.append(wilker)

        where_sql = " AND ".join(where_clauses)

        sql = f"""
            WITH threat_agg AS (
                SELECT
                    poly.id,
                    COUNT(DISTINCT obs.id) AS hotspot_count,
                    COUNT(DISTINCT obs.id) FILTER (WHERE obs.layer_key = 'perimeter_threat') AS non_kps_count,
                    COUNT(DISTINCT obs.id) FILTER (WHERE obs.layer_key != 'perimeter_threat') AS neighbor_kps_count,
                    ROUND(MIN(ST_Distance(poly.geometry::geography, obs.geom::geography)))::int AS min_dist_m
                FROM polygon_metadata poly
                JOIN hotspot_observations obs ON {where_sql}
                GROUP BY poly.id
            )
            SELECT
                COUNT(*) AS total_kps,
                COALESCE(SUM(hotspot_count), 0) AS total_hotspots,
                COALESCE(SUM(non_kps_count), 0) AS total_non_kps_hotspots,
                COALESCE(SUM(neighbor_kps_count), 0) AS total_neighbor_kps_hotspots,
                COUNT(*) FILTER (WHERE min_dist_m < 1000) AS bahaya_count,
                COUNT(*) FILTER (WHERE min_dist_m >= 1000 AND min_dist_m < 3000) AS waspada_count,
                COUNT(*) FILTER (WHERE min_dist_m >= 3000) AS pantau_count
            FROM threat_agg;
        """

        with self.postgres_store.connection() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, params)
                row = cur.fetchone()
                if not row:
                    res: dict[str, Any] = {
                        "total_kps_threatened": 0,
                        "bahaya_count": 0,
                        "waspada_count": 0,
                        "pantau_count": 0,
                        "total_external_hotspots": 0,
                        "non_kps_hotspots": 0,
                        "neighbor_kps_hotspots": 0,
                        "time_window_hours": time_window_hours,
                        "max_distance_km": max_distance_km,
                    }
                else:
                    res = {
                        "total_kps_threatened": int(row["total_kps"] or 0),
                        "total_external_hotspots": int(row["total_hotspots"] or 0),
                        "non_kps_hotspots": int(row["total_non_kps_hotspots"] or 0),
                        "neighbor_kps_hotspots": int(row["total_neighbor_kps_hotspots"] or 0),
                        "bahaya_count": int(row["bahaya_count"] or 0),
                        "waspada_count": int(row["waspada_count"] or 0),
                        "pantau_count": int(row["pantau_count"] or 0),
                        "time_window_hours": time_window_hours,
                        "max_distance_km": max_distance_km,
                    }
                self.cache_service.write(cache_key, res, ttl_hours=1)
                return res

    def get_threats(
        self,
        time_window_hours: int = 48,
        max_distance_km: float = 5.0,
        level: str | None = None,
        province: str | None = None,
        regency: str | None = None,
        wilker: str | None = None,
        search: str | None = None,
        limit: int = 100,
        offset: int = 0,
    ) -> dict[str, Any]:
        """Daftar KPS yang terancam api luar, diurutkan dari jarak terdekat."""
        cache_key = f"fire_spread_threats_{time_window_hours}_{max_distance_km:.2f}_{level or 'all'}_{province or 'all'}_{regency or 'all'}_{wilker or 'all'}_{search or 'none'}_{limit}_{offset}"
        cached = self.cache_service.read(cache_key)
        if cached is not None and isinstance(cached, dict):
            return cached

        if not self.postgres_store.enabled:
            return {"items": [], "total": 0, "limit": limit, "offset": offset}

        max_deg = max_distance_km * 1000.0 * DEGREE_PER_METER
        where_clauses = [
            "poly.is_active = true",
            f"obs.detected_at >= NOW() - INTERVAL '{int(time_window_hours)} hours'",
            f"ST_DWithin(poly.geometry, obs.geom, {max_deg:.6f})",
            "NOT ST_Intersects(poly.geometry, obs.geom)",
        ]
        params: list[Any] = []
        if province:
            where_clauses.append("poly.nama_prov = %s")
            params.append(province)
        if regency:
            where_clauses.append("poly.nama_kab = %s")
            params.append(regency)
        if wilker:
            where_clauses.append("poly.wilker_bps = %s")
            params.append(wilker)
        if search:
            where_clauses.append(
                "(poly.lembaga ILIKE %s OR poly.nama_desa ILIKE %s OR poly.nama_kab ILIKE %s)"
            )
            term = f"%{search.strip()}%"
            params.extend([term, term, term])

        where_sql = " AND ".join(where_clauses)

        sql = f"""
            WITH candidate_threats AS (
                SELECT
                    poly.id AS polygon_id,
                    poly.lembaga,
                    poly.lembaga AS nama_kps,
                    poly.nama_desa,
                    poly.nama_kec,
                    poly.nama_kab,
                    poly.nama_prov,
                    poly.skema,
                    poly.no_sk,
                    poly.wilker_bps,
                    poly.luas_final AS luas_ha,
                    poly.geometry,
                    COUNT(DISTINCT obs.id) AS external_hotspots_count,
                    COUNT(DISTINCT obs.id) FILTER (WHERE obs.layer_key = 'perimeter_threat') AS non_kps_count,
                    COUNT(DISTINCT obs.id) FILTER (WHERE obs.layer_key != 'perimeter_threat') AS neighbor_kps_count,
                    ROUND(MIN(ST_Distance(poly.geometry::geography, obs.geom::geography)))::int AS min_distance_m,
                    ROUND(MAX(COALESCE((obs.raw_payload->>'frp')::float, 0.0))::numeric, 1) AS max_frp,
                    ROUND(AVG(COALESCE((obs.raw_payload->>'frp')::float, 0.0))::numeric, 1) AS avg_frp
                FROM polygon_metadata poly
                JOIN hotspot_observations obs ON {where_sql}
                GROUP BY poly.id
            ),
            filtered_threats AS (
                SELECT *, COUNT(*) OVER() AS full_count
                FROM candidate_threats
                WHERE (%s::text IS NULL OR
                       (%s = 'bahaya' AND min_distance_m < 1000) OR
                       (%s = 'waspada' AND min_distance_m >= 1000 AND min_distance_m < 3000) OR
                       (%s = 'pantau' AND min_distance_m >= 3000))
                ORDER BY min_distance_m ASC, external_hotspots_count DESC
                LIMIT %s OFFSET %s
            )
            SELECT
                f.*,
                degrees(ST_Azimuth(ST_Centroid(f.geometry), nearest.geom)) AS bearing_deg,
                ST_AsGeoJSON(ST_ClosestPoint(f.geometry, nearest.geom)) AS nearest_boundary_point_geojson,
                ST_AsGeoJSON(nearest.geom) AS nearest_hotspot_geojson,
                nearest.satellite AS nearest_satellite,
                nearest.confidence AS nearest_confidence,
                nearest.layer_key AS nearest_layer_key,
                nearest.agency_name AS nearest_agency_name,
                to_char(nearest.detected_at, 'YYYY-MM-DD HH24:MI:SS OF') AS nearest_detected_at
            FROM filtered_threats f
            LEFT JOIN LATERAL (
                SELECT obs2.*
                FROM hotspot_observations obs2
                WHERE obs2.detected_at >= NOW() - INTERVAL '{int(time_window_hours)} hours'
                  AND ST_DWithin(f.geometry, obs2.geom, {max_deg:.6f})
                  AND NOT ST_Intersects(f.geometry, obs2.geom)
                ORDER BY ST_Distance(f.geometry::geography, obs2.geom::geography) ASC
                LIMIT 1
            ) nearest ON true;
        """

        exec_params = list(params) + [level, level, level, level, limit, offset]

        items = []
        total_count = 0
        with self.postgres_store.connection() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, exec_params)
                rows = cur.fetchall()
                for r in rows:
                    total_count = int(r["full_count"])
                    dist_m = int(r["min_distance_m"])
                    lvl, lvl_label = distance_to_level(dist_m)
                    bearing = float(r["bearing_deg"]) if r.get("bearing_deg") is not None else None
                    compass = degrees_to_compass(bearing)

                    near_pt = json.loads(r["nearest_boundary_point_geojson"]) if r.get("nearest_boundary_point_geojson") else None
                    near_hs = json.loads(r["nearest_hotspot_geojson"]) if r.get("nearest_hotspot_geojson") else None

                    if dist_m < 1000:
                        rekomendasi = f"DARURAT: Api berjarak {dist_m} m dari arah {compass}. Segera terjunkan tim patroli batas & buat sekat bakar darurat!"
                    elif dist_m < 3000:
                        rekomendasi = f"SIAGA: Api aktif {dist_m/1000:.1f} km di arah {compass}. Hubungi KPH/desa tetangga & siagakan posko karhutla."
                    else:
                        rekomendasi = f"PANTAU: Pantau pergerakan klaster api {dist_m/1000:.1f} km di arah {compass} via satelit secara berkala."

                    is_non_kps = (r.get("nearest_layer_key") == "perimeter_threat")
                    origin = "non_kps" if is_non_kps else "neighbor_kps"
                    origin_label = "Luar Kawasan (Bukan KPS)" if is_non_kps else f"KPS Tetangga ({r.get('nearest_agency_name') or 'Lain'})"

                    items.append({
                        "polygon_id": int(r["polygon_id"]),
                        "lembaga": r.get("lembaga") or "-",
                        "nama_kps": r.get("nama_kps") or r.get("lembaga") or "-",
                        "nama_desa": r.get("nama_desa"),
                        "nama_kec": r.get("nama_kec"),
                        "nama_kab": r.get("nama_kab"),
                        "nama_prov": r.get("nama_prov"),
                        "skema": r.get("skema"),
                        "no_sk": r.get("no_sk"),
                        "wilker_bps": r.get("wilker_bps"),
                        "luas_ha": float(r["luas_ha"]) if r.get("luas_ha") is not None else None,
                        "min_distance_m": dist_m,
                        "min_distance_km": round(dist_m / 1000.0, 2),
                        "status_level": lvl,
                        "status_label": lvl_label,
                        "external_hotspots_count": int(r["external_hotspots_count"]),
                        "non_kps_hotspots_count": int(r.get("non_kps_count") or 0),
                        "neighbor_kps_hotspots_count": int(r.get("neighbor_kps_count") or 0),
                        "threat_origin": origin,
                        "threat_origin_label": origin_label,
                        "max_frp": float(r["max_frp"]) if r.get("max_frp") is not None else 0.0,
                        "avg_frp": float(r["avg_frp"]) if r.get("avg_frp") is not None else 0.0,
                        "bearing_deg": round(bearing, 1) if bearing is not None else None,
                        "bearing_compass": compass,
                        "rekomendasi": rekomendasi,
                        "nearest_hotspot": {
                            "coordinates": near_hs.get("coordinates") if near_hs else None,
                            "satellite": r.get("nearest_satellite"),
                            "confidence": r.get("nearest_confidence"),
                            "detected_at": r.get("nearest_detected_at"),
                            "layer_key": r.get("nearest_layer_key"),
                            "agency_name": r.get("nearest_agency_name"),
                        },
                        "nearest_boundary_point": near_pt.get("coordinates") if near_pt else None,
                    })

        result = {
            "items": items,
            "total": total_count,
            "limit": limit,
            "offset": offset,
            "time_window_hours": time_window_hours,
            "max_distance_km": max_distance_km,
        }
        self.cache_service.write(cache_key, result, ttl_hours=1)
        return result

    def get_threat_detail(
        self,
        polygon_id: int,
        time_window_hours: int = 48,
        max_distance_km: float = 5.0,
    ) -> dict[str, Any] | None:
        """Detail satu KPS terancam beserta GeoJSON poligon dan seluruh titik api luar."""
        cache_key = f"fire_spread_detail_{polygon_id}_{time_window_hours}_{max_distance_km:.2f}"
        cached = self.cache_service.read(cache_key)
        if cached is not None and isinstance(cached, dict):
            return cached

        if not self.postgres_store.enabled:
            return None

        max_deg = max_distance_km * 1000.0 * DEGREE_PER_METER

        with self.postgres_store.connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT
                        id, lembaga, nama_desa, nama_kec, nama_kab, nama_prov,
                        skema, no_sk, wilker_bps, luas_final AS luas_ha,
                        ST_AsGeoJSON(geometry) AS geometry_json,
                        ST_AsGeoJSON(ST_Centroid(geometry)) AS centroid_json
                    FROM polygon_metadata
                    WHERE id = %s AND is_active = true
                    LIMIT 1;
                    """,
                    (polygon_id,),
                )
                poly = cur.fetchone()
                if not poly:
                    return None

                geom_json = json.loads(poly["geometry_json"]) if poly.get("geometry_json") else None
                centroid = json.loads(poly["centroid_json"]) if poly.get("centroid_json") else None

                # Query KPS bersebelahan/tetangga dalam radius buffer
                cur.execute(
                    f"""
                    SELECT
                        p2.id,
                        p2.lembaga,
                        p2.nama_desa,
                        p2.nama_kec,
                        p2.nama_kab,
                        p2.skema,
                        p2.luas_final AS luas_ha,
                        ROUND(ST_Distance(p_target.geometry::geography, p2.geometry::geography))::int AS distance_m,
                        ST_AsGeoJSON(p2.geometry) AS geometry_json,
                        (
                            SELECT COUNT(*)
                            FROM hotspot_observations h
                            WHERE h.detected_at >= NOW() - INTERVAL '{int(time_window_hours)} hours'
                              AND ST_Intersects(p2.geometry, h.geom)
                        ) AS hotspot_count
                    FROM polygon_metadata p_target
                    CROSS JOIN LATERAL (
                        SELECT id, lembaga, nama_desa, nama_kec, nama_kab, skema, luas_final, geometry
                        FROM polygon_metadata
                        WHERE is_active = true
                          AND id != p_target.id
                          AND geometry && ST_Expand(p_target.geometry, {max_deg:.6f})
                          AND ST_DWithin(p_target.geometry, geometry, {max_deg:.6f})
                    ) p2
                    WHERE p_target.id = %s
                    ORDER BY distance_m ASC
                    LIMIT 30;
                    """,
                    (polygon_id,),
                )
                neighbor_rows = cur.fetchall()
                neighbors = []
                for nr in neighbor_rows:
                    n_geom = json.loads(nr["geometry_json"]) if nr.get("geometry_json") else None
                    if n_geom:
                        neighbors.append({
                            "id": int(nr["id"]),
                            "lembaga": nr.get("lembaga") or "-",
                            "nama_desa": nr.get("nama_desa"),
                            "nama_kec": nr.get("nama_kec"),
                            "nama_kab": nr.get("nama_kab"),
                            "skema": nr.get("skema"),
                            "luas_ha": float(nr["luas_ha"]) if nr.get("luas_ha") is not None else None,
                            "distance_m": int(nr["distance_m"]),
                            "distance_km": round(int(nr["distance_m"]) / 1000.0, 2),
                            "hotspot_count": int(nr["hotspot_count"] or 0),
                            "geometry": n_geom,
                        })

                # Query hotspot baik di perimeter luar maupun di dalam kawasan
                cur.execute(
                    f"""
                    SELECT
                        obs.id,
                        obs.latitude,
                        obs.longitude,
                        obs.satellite,
                        obs.confidence,
                        obs.brightness,
                        COALESCE((obs.raw_payload->>'frp')::float, 0.0) AS frp,
                        to_char(obs.detected_at, 'YYYY-MM-DD HH24:MI:SS OF') AS detected_at_str,
                        ST_Intersects(poly.geometry, obs.geom) AS is_inside,
                        CASE
                            WHEN ST_Intersects(poly.geometry, obs.geom) THEN 0
                            ELSE ROUND(ST_Distance(poly.geometry::geography, obs.geom::geography))::int
                        END AS distance_m,
                        degrees(ST_Azimuth(ST_Centroid(poly.geometry), obs.geom)) AS bearing_deg,
                        CASE
                            WHEN ST_Intersects(poly.geometry, obs.geom) THEN NULL
                            ELSE ST_AsGeoJSON(ST_ClosestPoint(poly.geometry, obs.geom))
                        END AS closest_pt_geojson
                    FROM hotspot_observations obs
                    JOIN polygon_metadata poly ON poly.id = %s
                    WHERE obs.detected_at >= NOW() - INTERVAL '{int(time_window_hours)} hours'
                      AND ST_DWithin(poly.geometry, obs.geom, {max_deg:.6f})
                    ORDER BY is_inside DESC, distance_m ASC;
                    """,
                    (polygon_id,),
                )
                hs_rows = cur.fetchall()

                hotspots = []
                closest_vector = None
                closest_ext_found = False
                for h in hs_rows:
                    is_inside = bool(h["is_inside"])
                    dist_m = int(h["distance_m"])
                    bearing = float(h["bearing_deg"]) if h.get("bearing_deg") is not None else None
                    c_pt = json.loads(h["closest_pt_geojson"]) if h.get("closest_pt_geojson") else None

                    is_non_kps = (h.get("layer_key") == "perimeter_threat")
                    if is_inside:
                        lvl = "internal"
                        lvl_label = "Di Dalam Kawasan"
                        origin = "internal"
                        origin_label = "Di Dalam Kawasan"
                    else:
                        lvl, lvl_label = distance_to_level(dist_m)
                        origin = "non_kps" if is_non_kps else "neighbor_kps"
                        origin_label = "Luar Kawasan (Bukan KPS)" if is_non_kps else f"KPS Tetangga ({h.get('agency_name') or 'Lain'})"

                    compass = degrees_to_compass(bearing) if not is_inside else "Dalam Kawasan"

                    item = {
                        "id": h["id"],
                        "latitude": float(h["latitude"]),
                        "longitude": float(h["longitude"]),
                        "satellite": h.get("satellite"),
                        "confidence": h.get("confidence"),
                        "brightness": float(h["brightness"]) if h.get("brightness") is not None else None,
                        "frp": float(h["frp"]) if h.get("frp") is not None else 0.0,
                        "detected_at": h.get("detected_at_str"),
                        "is_inside": is_inside,
                        "layer_key": h.get("layer_key"),
                        "agency_name": h.get("agency_name"),
                        "threat_origin": origin,
                        "threat_origin_label": origin_label,
                        "distance_m": dist_m,
                        "distance_km": round(dist_m / 1000.0, 2),
                        "status_level": lvl,
                        "status_label": lvl_label,
                        "bearing_deg": round(bearing, 1) if bearing is not None else None,
                        "bearing_compass": compass,
                        "closest_kps_point": c_pt.get("coordinates") if c_pt else None,
                    }
                    hotspots.append(item)
                    if not is_inside and not closest_ext_found and c_pt:
                        closest_vector = {
                            "hotspot_coords": [float(h["longitude"]), float(h["latitude"])],
                            "kps_boundary_coords": c_pt.get("coordinates"),
                            "distance_m": dist_m,
                            "bearing_compass": item["bearing_compass"],
                        }
                        closest_ext_found = True

                ext_hotspots = [h for h in hotspots if not h["is_inside"]]
                int_hotspots = [h for h in hotspots if h["is_inside"]]
                min_dist = ext_hotspots[0]["distance_m"] if ext_hotspots else 0
                status_lvl, status_lbl = distance_to_level(min_dist)

                res_detail = {
                    "polygon_id": int(poly["id"]),
                    "lembaga": poly.get("lembaga") or "-",
                    "nama_kps": poly.get("nama_kps") or poly.get("lembaga") or "-",
                    "nama_desa": poly.get("nama_desa"),
                    "nama_kec": poly.get("nama_kec"),
                    "nama_kab": poly.get("nama_kab"),
                    "nama_prov": poly.get("nama_prov"),
                    "skema": poly.get("skema"),
                    "no_sk": poly.get("no_sk"),
                    "wilker_bps": poly.get("wilker_bps"),
                    "luas_ha": float(poly["luas_ha"]) if poly.get("luas_ha") is not None else None,
                    "geometry": geom_json,
                    "centroid": centroid.get("coordinates") if centroid else None,
                    "status_level": status_lvl,
                    "status_label": status_lbl,
                    "min_distance_m": min_dist,
                    "min_distance_km": round(min_dist / 1000.0, 2),
                    "total_external_hotspots": len(ext_hotspots),
                    "total_non_kps_hotspots": len([h for h in ext_hotspots if h.get("threat_origin") == "non_kps"]),
                    "total_neighbor_kps_hotspots": len([h for h in ext_hotspots if h.get("threat_origin") == "neighbor_kps"]),
                    "total_internal_hotspots": len(int_hotspots),
                    "hotspots": hotspots,
                    "neighbors": neighbors,
                    "closest_vector": closest_vector,
                    "time_window_hours": time_window_hours,
                    "max_distance_km": max_distance_km,
                }
                self.cache_service.write(cache_key, res_detail, ttl_hours=1)
                return res_detail

    def export_threats_xlsx(
        self,
        time_window_hours: int = 48,
        max_distance_km: float = 5.0,
        level: str | None = None,
        province: str | None = None,
        regency: str | None = None,
        wilker: str | None = None,
        search: str | None = None,
    ) -> bytes:
        """Menghasilkan file Excel (.xlsx) laporan KPS terancam api luar."""
        data = self.get_threats(
            time_window_hours=time_window_hours,
            max_distance_km=max_distance_km,
            level=level,
            province=province,
            regency=regency,
            wilker=wilker,
            search=search,
            limit=5000,
            offset=0,
        )
        items = data.get("items", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Siaga Rambatan Api"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Arial", size=14, bold=True, color="1F2937")
        font_sub = Font(name="Arial", size=9, italic=True, color="4B5563")
        font_th = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        font_cell = Font(name="Arial", size=9, color="111827")
        font_bold = Font(name="Arial", size=9, bold=True, color="111827")

        fill_th = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        fill_orange = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
        fill_yellow = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

        thin = Side(border_style="thin", color="D1D5DB")
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws["A1"] = "LAPORAN SIAGA RAMBATAN API (DETEKSI ANCAMAN LUAR KPS)"
        ws["A1"].font = font_title
        ws["A2"] = (
            f"Waktu Pantau: {time_window_hours} Jam Terakhir | Radius Buffer: {max_distance_km} km | "
            f"Diunduh pada: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | "
            f"Total KPS Terancam: {len(items)}"
        )
        ws["A2"].font = font_sub

        headers = [
            "No", "Status Siaga", "Jarak Terdekat", "Arah Ancaman", "Jml Titik Api Luar",
            "Nama KPS / Lembaga", "Skema", "Desa", "Kecamatan", "Kabupaten", "Provinsi",
            "Wilker BPS", "Max FRP (MW)", "Waktu Deteksi Terdekat", "Rekomendasi Aksi Lapangan"
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = font_th
            cell.fill = fill_th
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_all
        ws.row_dimensions[row_idx].height = 28

        for idx, item in enumerate(items, 1):
            row_idx += 1
            lvl = item["status_level"]
            fill_lvl = fill_red if lvl == "bahaya" else (fill_orange if lvl == "waspada" else fill_yellow)

            row_data = [
                idx,
                item["status_label"],
                f"{item['min_distance_m']} m ({item['min_distance_km']} km)",
                f"{item['bearing_compass']} ({item['bearing_deg']}°)" if item['bearing_deg'] is not None else "-",
                item["external_hotspots_count"],
                item["lembaga"],
                item.get("skema") or "-",
                item.get("nama_desa") or "-",
                item.get("nama_kec") or "-",
                item.get("nama_kab") or "-",
                item.get("nama_prov") or "-",
                item.get("wilker_bps") or "-",
                item["max_frp"],
                item.get("nearest_hotspot", {}).get("detected_at") or "-",
                item["rekomendasi"]
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_cell
                cell.border = border_all
                if col_idx in (1, 3, 4, 5, 7, 13):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.fill = fill_lvl
                    cell.font = font_bold
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 22

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

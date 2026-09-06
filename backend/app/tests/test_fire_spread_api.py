import io
import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.fire_spread import router
from app.services.fire_spread_service import (
    FireSpreadService,
    degrees_to_compass,
    distance_to_level,
)


def test_degrees_to_compass():
    assert degrees_to_compass(0) == "Utara (N)"
    assert degrees_to_compass(45) == "Timur Laut (NE)"
    assert degrees_to_compass(90) == "Timur (E)"
    assert degrees_to_compass(135) == "Tenggara (SE)"
    assert degrees_to_compass(180) == "Selatan (S)"
    assert degrees_to_compass(225) == "Barat Daya (SW)"
    assert degrees_to_compass(270) == "Barat (W)"
    assert degrees_to_compass(315) == "Barat Laut (NW)"
    assert degrees_to_compass(350) == "Utara (N)"
    assert degrees_to_compass(None) == "Tidak diketahui"


def test_distance_to_level():
    assert distance_to_level(500) == ("bahaya", "Bahaya Kritis (< 1 km)")
    assert distance_to_level(1500) == ("waspada", "Waspada (1–3 km)")
    assert distance_to_level(4500) == ("pantau", "Pantau (3–5 km)")


class _FakePostgresStore:
    enabled = False


def _build_client():
    app = FastAPI()
    app.include_router(router, prefix="/api")
    return TestClient(app)


def test_fire_spread_summary_endpoint_disabled_store(monkeypatch):
    client = _build_client()
    monkeypatch.setattr(
        "app.api.fire_spread.FireSpreadService",
        lambda *a, **k: FireSpreadService(postgres_store=_FakePostgresStore()),
    )
    resp = client.get("/api/fire-spread/summary?time_window_hours=48&max_distance_km=5")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_kps_threatened"] == 0
    assert data["bahaya_count"] == 0
    assert data["waspada_count"] == 0
    assert data["pantau_count"] == 0


def test_fire_spread_threats_endpoint_mock(monkeypatch):
    client = _build_client()

    class _MockService:
        def get_threats(self, **kwargs):
            return {
                "items": [
                    {
                        "polygon_id": 101,
                        "lembaga": "KPS Damai",
                        "nama_kps": "KPS Damai",
                        "nama_kab": "Lamandau",
                        "nama_prov": "Kalimantan Tengah",
                        "min_distance_m": 450,
                        "min_distance_km": 0.45,
                        "status_level": "bahaya",
                        "status_label": "Bahaya Kritis (< 1 km)",
                        "external_hotspots_count": 12,
                        "max_frp": 35.5,
                        "bearing_deg": 315.0,
                        "bearing_compass": "Barat Laut (NW)",
                        "rekomendasi": "DARURAT: Api berjarak 450 m dari arah Barat Laut (NW)",
                    }
                ],
                "total": 1,
                "limit": 100,
                "offset": 0,
            }

    monkeypatch.setattr("app.api.fire_spread.FireSpreadService", _MockService)
    resp = client.get("/api/fire-spread/threats?level=bahaya")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] == 1
    assert data["items"][0]["status_level"] == "bahaya"
    assert data["items"][0]["bearing_compass"] == "Barat Laut (NW)"


def test_fire_spread_export_excel_mock(monkeypatch):
    client = _build_client()

    class _MockService:
        def export_threats_xlsx(self, **kwargs):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Siaga Rambatan Api"
            ws["A1"] = "Test Export"
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

    monkeypatch.setattr("app.api.fire_spread.FireSpreadService", _MockService)
    resp = client.get("/api/fire-spread/export.xlsx")
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Siaga Rambatan Api" in wb.sheetnames


def test_fire_spread_detail_endpoint_mock(monkeypatch):
    client = _build_client()

    class _MockService:
        def get_threat_detail(self, polygon_id, time_window_hours=48, max_distance_km=5.0):
            if polygon_id == 999999:
                return None
            return {
                "polygon_id": polygon_id,
                "lembaga": "LPHD PEMATANG GADUNG",
                "nama_kps": "LPHD PEMATANG GADUNG",
                "geometry": {"type": "Polygon", "coordinates": [[[109.9, -1.8], [110.0, -1.8], [110.0, -1.9], [109.9, -1.9], [109.9, -1.8]]]},
                "centroid": [109.95, -1.85],
                "status_level": "bahaya",
                "status_label": "Bahaya Kritis (< 1 km)",
                "min_distance_m": 59,
                "min_distance_km": 0.06,
                "total_external_hotspots": 152,
                "total_internal_hotspots": 33,
                "hotspots": [
                    {
                        "id": 1,
                        "latitude": -1.85,
                        "longitude": 109.95,
                        "is_inside": True,
                        "status_level": "internal",
                        "status_label": "Di Dalam Kawasan",
                        "distance_m": 0,
                    },
                    {
                        "id": 2,
                        "latitude": -1.799,
                        "longitude": 109.95,
                        "is_inside": False,
                        "status_level": "bahaya",
                        "status_label": "Bahaya Kritis (< 1 km)",
                        "distance_m": 59,
                    },
                ],
                "neighbors": [
                    {
                        "id": 287888,
                        "lembaga": "LPHD SUNGAI BESAR",
                        "distance_m": 0,
                        "distance_km": 0.0,
                        "hotspot_count": 139,
                        "geometry": {"type": "Polygon", "coordinates": [[[109.9, -1.7], [110.0, -1.7], [110.0, -1.8], [109.9, -1.8], [109.9, -1.7]]]},
                    }
                ],
                "closest_vector": None,
                "time_window_hours": time_window_hours,
                "max_distance_km": max_distance_km,
            }

    monkeypatch.setattr("app.api.fire_spread.FireSpreadService", _MockService)
    resp = client.get("/api/fire-spread/detail?polygon_id=101")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_internal_hotspots"] == 33
    assert data["total_external_hotspots"] == 152
    assert len(data["neighbors"]) == 1
    assert data["neighbors"][0]["lembaga"] == "LPHD SUNGAI BESAR"
    assert data["hotspots"][0]["is_inside"] is True

    # Test 404
    resp_404 = client.get("/api/fire-spread/detail?polygon_id=999999")
    assert resp_404.status_code == 404


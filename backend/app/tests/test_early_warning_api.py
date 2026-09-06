import io
import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.early_warning import router
from app.core.auth import issue_token
from app.core.session_store import get_auth_store


class _DummyAuthStore:
    enabled = False


def _build_client():
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_auth_store] = lambda: _DummyAuthStore()
    return TestClient(app)


def test_export_early_warning_excel_post_with_items():
    client = _build_client()
    token = issue_token(user_id=1, username="admin", role="admin")

    items = [
        {
            "id": 101,
            "lembaga": "KPS Hutan Lestari",
            "wilker_bps": "BPSKL Jawa",
            "skema": "HD",
            "nama_desa": "Desa Makmur",
            "nama_kec": "Kecamatan Subur",
            "nama_kab": "Kabupaten Hijau",
            "nama_prov": "Jawa Barat",
            "luas_sk": 250.0,
            "total_burned_ha": 12.5,
            "burn_frequency": 2,
            "hotspots_today": 3,
            "hotspots_today_strict_reburn": 1,
            "hotspots_today_expanding": 2,
            "min_distance_km": 0.5,
            "max_distance_km": 1.2,
            "avg_distance_km": 0.85,
            "fire_direction": "Timur (T)",
            "fire_azimuth_deg": 90.0,
            "propagation_zone": "Kombinasi Bara & Merambat",
            "zone_code": "combo",
            "hotspots_yesterday": 0,
            "hotspots_7d": 5,
            "hotspots_month": 8,
            "hotspots_year": 12,
            "latest_hotspot_at": "2026-09-06T10:00:00Z",
            "ftri_score": 78.5,
            "status_label": "Strict Re-burn & Blok Baru",
        }
    ]

    response = client.post(
        "/api/early-warning/export.xlsx",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "title": "REKAP FILTERED TEST",
            "subtitle": "Kategori: Ada Hotspot Hari Ini | Wilayah: Jawa Barat",
            "items": items,
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws["A1"].value == "REKAP FILTERED TEST"
    assert "Jawa Barat" in str(ws["A2"].value)
    assert ws.cell(row=5, column=3).value == "KPS Hutan Lestari"
    assert ws.cell(row=5, column=10).value == 250.0
    assert ws.cell(row=5, column=13).value == 3
